from __future__ import annotations
import re
from pathlib import Path
from collections import OrderedDict
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path.cwd()
OUT = ROOT / 'outputs' / 'norzagapay-data-dictionary-20261022' / 'NorzAgapay_Data_Dictionary.docx'
SCHEMA = ROOT / 'database' / 'schema.sql'
MIGRATIONS = sorted((ROOT / 'database' / 'migrations').glob('*.sql'))

# Limit the dictionary to tables that carry the main web incident cycle.
GROUPS = [
  ('Accounts and barangay context', 'These tables identify report submitters, response staff, and barangay jurisdictions.', [
    ('resident_user', 'Resident account used to identify people who submit and track reports.'),
    ('users', 'MDRRMO staff and responder accounts used for dispatch and field actions.'),
    ('barangay_users', 'Barangay staff and responder accounts used for local assignments.'),
    ('barangays', 'Barangay names and locations used to route reports and coordinate response.'),
  ]),
  ('Incident report records', 'These tables store reports from submission through review and resolution.', [
    ('barangay_reports', 'Barangay incident record used for local response updates and resolution.'),
    ('mdrrmo_reports', 'MDRRMO report record with incident details, evidence references, status, and resolution data.'),
  ]),
  ('Responder units and rosters', 'These tables define available MDRRMO teams and the people assigned to each unit.', [
    ('officers', 'Responder qualifications and officer details linked to MDRRMO unit rosters.'),
    ('respond_units', 'MDRRMO response-unit name, specialization, availability, and activation state.'),
    ('respond_unit_members', 'Officer and responder membership and roles within each response unit.'),
  ]),
  ('Dispatch, assignments, and crews', 'These tables link reports to MDRRMO or barangay responders and record selected crews.', [
    ('mdrrmo_report_assignments', 'Links an MDRRMO report to assigned responders and tracks their response state.'),
    ('barangay_report_assignments', 'Links a barangay report to barangay responders and their response roles.'),
    ('mdrrmo_report_barangay_assignments', 'Tracks MDRRMO-to-barangay dispatch, arrival, response, and resolution.'),
    ('mdrrmo_assignment_crew_members', 'Records the MDRRMO crew selected for a responder assignment.'),
    ('mdrrmo_report_barangay_crew_members', 'Records barangay members selected for a coordinated response.'),
  ]),
  ('Status and assignment history', 'These tables preserve report status changes and barangay handoff events.', [
    ('mdrrmo_report_lifecycle_history', 'Records the MDRRMO report status and action history.'),
    ('barangay_report_lifecycle_history', 'Records the barangay report status and action history.'),
    ('mdrrmo_report_barangay_assignment_events', 'Records barangay assignment, reassignment, recall, and coordination events.'),
  ]),
  ('Assistance and transport', 'These tables support requests for extra help and injured-person transport.', [
    ('barangay_assistance_requests', 'Records a barangay request for responders, equipment, resources, or MDRRMO help.'),
    ('resource_requests', 'Records resource requests submitted for dispatcher review.'),
    ('hospital_directory', 'Stores hospital names and coordinates used to select nearby transport destinations.'),
  ]),
]

# Enumerated PostgreSQL types are documented separately so field rows stay readable.
ENUM_VALUES = OrderedDict([
  ('user_role', 'master_admin, admin, logistics, dispatcher, responder'),
  ('unit_type', 'police, fire, medical, Rescue Officer, Swift Water Rescue Officer, Mountain Rescue Officer, Emergency Medical Responder (EMR), Ambulance Officer / EMS Personnel, Fire Response Officer, Evacuation Officer, Safety & Security Officer, Traffic & Road Clearing Officer, Communications Officer, Logistics Response Officer, Damage Assessment Officer'),
  ('user_status', 'active, inactive, pending_verification, occupied, rejected'),
  ('report_status', 'pending, verified, rejected, responding, escalated, resolved'),
  ('alert_level', 'normal, warning, critical'),
  ('alert_type', 'weather, flood, earthquake, dam'),
  ('alert_status', 'active, acknowledged, resolved'),
  ('risk_level', 'low, moderate, high, critical'),
  ('barangay_role', 'admin, dispatcher, responder, staff'),
])


def strip_sql_comment(line: str) -> str:
    # SQL comments in the schema are trailing and do not contain quoted '--'.
    return line.split('--', 1)[0].strip().rstrip(',').strip()


def parse_column(name: str, raw: str) -> dict:
    raw = strip_sql_comment(raw).strip().rstrip(',')
    cut = re.split(r'\s+(?=(?:NOT\s+NULL|NULL\b|DEFAULT\b|PRIMARY\s+KEY|UNIQUE\b|REFERENCES\b|CHECK\s*\(|COLLATE\b|GENERATED\b|CONSTRAINT\b))', raw, maxsplit=1, flags=re.I)
    typ = cut[0].strip().rstrip(',') if cut else raw
    typ = re.sub(r'\s+', ' ', typ)
    default = '—'
    dm = re.search(r'\bDEFAULT\s+(.+?)(?=\s+(?:NOT\s+NULL|PRIMARY\s+KEY|UNIQUE\b|REFERENCES\b|CHECK\s*\()|$)', raw, re.I)
    if dm:
        default = re.sub(r'\s+', ' ', dm.group(1).strip().rstrip(',')).strip()
    default = normalize_default(default)
    nullable = not bool(re.search(r'\bNOT\s+NULL\b|\bPRIMARY\s+KEY\b', raw, re.I))
    key = []
    if re.search(r'\bPRIMARY\s+KEY\b', raw, re.I):
        key.append('PK')
    if re.search(r'\bUNIQUE\b', raw, re.I):
        key.append('UQ')
    ref = re.search(r'\bREFERENCES\s+(?:public\.)?(\w+)\s*\((\w+)\)', raw, re.I)
    if ref:
        key.append(f'FK {ref.group(1)}.{ref.group(2)}')
    if re.search(r'\bCHECK\s*\(', raw, re.I):
        key.append('CHECK')
    return {'name': name, 'raw': raw, 'type': typ, 'nullable': nullable, 'default': default, 'keys': key}


def normalize_default(v: str) -> str:
    if not v or v == '—':
        return '—'
    v = re.sub(r'\s+', ' ', v.strip().rstrip(',')).strip()
    replacements = {
      "'[]'::jsonb": '[]', "'{}'::jsonb": '{}', "'[]'::json": '[]', "'{}'::json": '{}',
      "'{}'::text[]": '{}', "ARRAY[]::text[]": '{}', "ARRAY[]::uuid[]": '{}',
      "'[]'::text[]": '{}', "'{}'::uuid[]": '{}',
    }
    return replacements.get(v, v)


def split_table_body(body: str):
    fields = OrderedDict()
    for line in body.splitlines():
        clean = strip_sql_comment(line)
        if not clean:
            continue
        m = re.match(r'^"?([A-Za-z_][\w]*)"?\s+(.+)$', clean)
        if not m:
            continue
        name, raw = m.group(1), m.group(2)
        if name.upper() in {'CONSTRAINT', 'PRIMARY', 'UNIQUE', 'FOREIGN', 'CHECK', 'EXCLUDE'}:
            continue
        fields[name] = parse_column(name, raw)
    return fields


def split_constraints(body: str, fields: OrderedDict):
    for m in re.finditer(r'\bPRIMARY\s+KEY\s*\(([^)]+)\)', body, re.I):
        cols = re.findall(r'"?([A-Za-z_]\w*)"?', m.group(1))
        label = 'PK composite' if len(cols) > 1 else 'PK'
        for c in cols:
            if c in fields:
                fields[c]['nullable'] = False
                add_key(fields[c], label)
    for m in re.finditer(r'\bUNIQUE\s*\(([^)]+)\)', body, re.I):
        cols = re.findall(r'"?([A-Za-z_]\w*)"?', m.group(1))
        label = 'UQ composite' if len(cols) > 1 else 'UQ'
        for c in cols:
            if c in fields:
                add_key(fields[c], label)
    for m in re.finditer(r'\bFOREIGN\s+KEY\s*\(([^)]+)\)\s*REFERENCES\s+(?:public\.)?(\w+)\s*\(([^)]+)\)', body, re.I):
        cols = re.findall(r'"?([A-Za-z_]\w*)"?', m.group(1))
        refs = re.findall(r'"?([A-Za-z_]\w*)"?', m.group(3))
        for i, c in enumerate(cols):
            if c in fields:
                ref_col = refs[min(i, len(refs)-1)] if refs else 'id'
                add_key(fields[c], f'FK {m.group(2)}.{ref_col}')


def add_key(field, value):
    if value not in field['keys']:
        field['keys'].append(value)


def read_schema():
    sources = [(SCHEMA, SCHEMA.read_text(encoding='utf-8'))]
    sources += [(p, p.read_text(encoding='utf-8', errors='ignore')) for p in MIGRATIONS]
    tables = OrderedDict()
    create_re = re.compile(r'CREATE\s+TABLE\s+IF\s+NOT\s+EXISTS\s+(?:public\.)?([\w]+)\s*\((.*?)\n\);', re.I | re.S)
    # schema.sql is the consolidated definition; migrations supply only missing tables and deltas.
    for path, text in sources:
        for m in create_re.finditer(text):
            name, body = m.group(1).lower(), m.group(2)
            if name not in tables or path != SCHEMA:
                if name not in tables or path == SCHEMA:
                    fields = split_table_body(body)
                    split_constraints(body, fields)
                    tables[name] = {'fields': fields, 'body': body, 'source': path}
    # Apply declared ALTER TABLE deltas in schema and migration order.
    for path, text in sources:
        for am in re.finditer(r'ALTER\s+TABLE\s+(?:public\.)?([\w]+)(.*?);', text, re.I | re.S):
            name, stmt = am.group(1).lower(), am.group(2)
            if name not in tables:
                continue
            fields = tables[name]['fields']
            for drop in re.finditer(r'\bDROP\s+CONSTRAINT(?:\s+IF\s+EXISTS)?\s+"?([A-Za-z_]\w*)', stmt, re.I):
                constraint = drop.group(1).lower()
                prefix = name.lower() + '_'
                suffix = '_fkey'
                if constraint.startswith(prefix) and constraint.endswith(suffix):
                    column = constraint[len(prefix):-len(suffix)]
                    if column in fields:
                        fields[column]['keys'] = [key for key in fields[column]['keys'] if not key.upper().startswith('FK ')]
            for add in re.finditer(r'\bADD\s+COLUMN(?:\s+IF\s+NOT\s+EXISTS)?\s+"?([A-Za-z_]\w*)"?\s+(.+?)(?=\s*,?\s*ADD\s+COLUMN\b|$)', stmt, re.I | re.S):
                col, raw = add.group(1), add.group(2).strip().rstrip(',')
                if col not in fields:
                    fields[col] = parse_column(col, raw)
            for cm in re.finditer(r'\bALTER\s+COLUMN\s+"?([A-Za-z_]\w*)"?\s+(SET\s+DEFAULT\s+(.+?)|DROP\s+DEFAULT|SET\s+NOT\s+NULL|DROP\s+NOT\s+NULL|TYPE\s+([\w\s\[\](),]+?))(?=\s*,\s*ALTER\s+COLUMN\b|\s*,\s*ADD\s+COLUMN\b|$)', stmt, re.I | re.S):
                col = cm.group(1)
                if col not in fields:
                    continue
                action = cm.group(2).strip()
                if action.upper().startswith('SET DEFAULT'):
                    fields[col]['default'] = normalize_default(cm.group(3).strip())
                elif action.upper() == 'DROP DEFAULT':
                    fields[col]['default'] = '—'
                elif action.upper() == 'SET NOT NULL':
                    fields[col]['nullable'] = False
                elif action.upper() == 'DROP NOT NULL':
                    fields[col]['nullable'] = True
                elif action.upper().startswith('TYPE') and cm.group(4):
                    fields[col]['type'] = re.sub(r'\s+', ' ', cm.group(4).strip())
            # ALTER TABLE ... ADD CONSTRAINT ... key changes.
            for fm in re.finditer(r'\bFOREIGN\s+KEY\s*\(([^)]+)\)\s*REFERENCES\s+(?:public\.)?(\w+)\s*\(([^)]+)\)', stmt, re.I):
                cols = re.findall(r'"?([A-Za-z_]\w*)"?', fm.group(1))
                refs = re.findall(r'"?([A-Za-z_]\w*)"?', fm.group(3))
                for i, c in enumerate(cols):
                    if c in fields:
                        add_key(fields[c], f'FK {fm.group(2)}.{refs[min(i, len(refs)-1)] if refs else "id"}')
            for um in re.finditer(r'\bUNIQUE\s*\(([^)]+)\)', stmt, re.I):
                cols = re.findall(r'"?([A-Za-z_]\w*)"?', um.group(1))
                label = 'UQ composite' if len(cols) > 1 else 'UQ'
                for c in cols:
                    if c in fields:
                        add_key(fields[c], label)
    # Include standalone unique indexes declared in SQL. Partial indexes are
    # shown only when their predicate is IS NOT NULL; conditional role indexes
    # do not make a column globally unique.
    for path, text in sources:
        for im in re.finditer(r'CREATE\s+UNIQUE\s+INDEX\b[^;]*;', text, re.I | re.S):
            stmt = im.group(0)
            tm = re.search(r'\bON\s+(?:public\.)?(\w+)\s*(?:USING\s+\w+\s*)?\(', stmt, re.I)
            if not tm:
                continue
            name = tm.group(1).lower()
            where = re.search(r'\bWHERE\b(.+?);?$', stmt, re.I | re.S)
            if where and not re.fullmatch(r'\s*\w+\s+IS\s+NOT\s+NULL\s*;?\s*', where.group(1), re.I):
                continue
            start = tm.end()
            depth, end = 1, start
            while end < len(stmt) and depth:
                if stmt[end] == '(':
                    depth += 1
                elif stmt[end] == ')':
                    depth -= 1
                end += 1
            expr = stmt[start:end-1]
            expr = re.sub(r'\b(?:lower|upper)\s*\(\s*"?([A-Za-z_]\w*)"?\s*\)', r'\1', expr, re.I)
            cols = []
            for part in expr.split(','):
                part = part.strip().strip('"')
                if re.fullmatch(r'[A-Za-z_]\w*', part):
                    cols.append(part)
            if name in tables and cols:
                label = 'UQ composite' if len(cols) > 1 else 'UQ'
                for c in cols:
                    if c in tables[name]['fields']:
                        add_key(tables[name]['fields'][c], label)
    return tables


def type_and_length(raw_type: str):
    t = re.sub(r'\s+', ' ', raw_type.strip()).lower()
    length = '—'
    m = re.match(r'^(?:character varying|varchar|character|char)\s*\((\d+)\)', t)
    if m:
        return ('varchar' if 'varying' in t or 'varchar' in t else 'char', m.group(1))
    m = re.match(r'^(?:numeric|decimal)\s*\((\d+)(?:\s*,\s*(\d+))?\)', t)
    if m:
        return ('numeric', m.group(1) + (',' + m.group(2) if m.group(2) else ''))
    aliases = {
      'character varying': 'varchar', 'timestamp with time zone': 'timestamptz',
      'timestamp without time zone': 'timestamp', 'double precision': 'double precision',
      'boolean': 'boolean', 'int': 'integer', 'int4': 'integer', 'int8': 'bigint',
      'float8': 'double precision', 'float4': 'real', 'serial8': 'bigserial',
    }
    for prefix in sorted(aliases, key=len, reverse=True):
        if t == prefix or t.startswith(prefix + ' '):
            t = aliases[prefix] + t[len(prefix):]
            break
    return t, length


def compact_default(v: str):
    if len(v) > 32:
        return v[:29] + '...'
    return v


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement('w:cantSplit')
    trPr.append(cant)


def set_run_font(run, size=12, bold=None, italic=None, color=None):
    run.font.name = 'Times New Roman'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = RGBColor(*color)


def format_paragraph(p, before=0, after=3, line=1.0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def add_text(doc, text, style=None, bold=False, italic=False, align=None, before=0, after=3):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    format_paragraph(p, before, after)
    r = p.add_run(text)
    set_run_font(r, 12, bold, italic)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout'); tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, '3F6B8F')
        set_cell_margins(cell)
        p = cell.paragraphs[0]; format_paragraph(p, 0, 0, 1.0)
        r = p.add_run(head); set_run_font(r, 12, True, color=(255,255,255))
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, vals in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(vals):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2 == 1: set_cell_shading(cell, 'EEF3F7')
            set_cell_margins(cell)
            p = cell.paragraphs[0]; format_paragraph(p, 0, 0, 1.0)
            r = p.add_run(str(value)); set_run_font(r, 12, False, color=(25,35,45))
    return table


def key_text(keys):
    return '; '.join(keys) if keys else '—'


def main():
    tables = read_schema()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    # Keep every style and table cell in the requested font and size.
    for stylename in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        st = doc.styles[stylename]
        st.font.name = 'Times New Roman'
        st._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
        st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        st.font.size = Pt(12)
        st.font.color.rgb = RGBColor(20, 35, 50)
    doc.styles['Normal'].paragraph_format.space_after = Pt(3)
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(3)
    doc.styles['Heading 2'].font.bold = True
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
    doc.styles['Heading 2'].paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    format_paragraph(p, 0, 5)
    r = p.add_run('NorzAgapay Data Dictionary'); set_run_font(r, 12, True, color=(20,35,50))
    add_text(doc, 'Core data behind report submission, barangay coordination, MDRRMO dispatch, field response, and incident resolution.', italic=True, after=5)
    add_text(doc, 'Prepared 22 October 2026', after=6)
    add_text(doc, 'This dictionary focuses on the key records in the NorzAgapay web incident cycle: submit a report, review and assign it, record field response, and close and track the incident. Standalone weather monitoring, public broadcasts, and internal notification queues are outside this scope. Table definitions are based on the checked-in schema and migrations.', after=5)
    add_text(doc, 'Report records hold evidence references, response notes, and resolution details; uploaded media files themselves are stored in Supabase Storage. Assignment and history tables preserve who responded and how the incident progressed.', after=7)

    add_text(doc, 'Table Groups', style='Heading 1', after=3)
    summary = []
    for group, purpose, specs in GROUPS:
        summary.append((group, str(len(specs)), purpose))
    summary_table = add_table(doc, ['System Area', 'Tables', 'Purpose'], summary, [1.7, 0.7, 3.87])
    for row in summary_table.rows:
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(doc, 'Reading the field tables', style='Heading 1', before=7, after=3)
    add_text(doc, 'Allow Null is Yes when the SQL column permits NULL and No when it is required or is a primary key. A dash means that no explicit length, default, or key is declared. PK means primary key, FK identifies the referenced table and column, UQ means unique, and CHECK means the column has a database check constraint. PostgreSQL enum values are listed in the final section.', after=3)
    add_text(doc, 'Source files: database/schema.sql and database/migrations/*.sql; table usage was checked against the web dashboard and its backend.', italic=True, before=5, after=2)

    definitions_heading = add_text(doc, 'Detailed table definitions', style='Heading 1', after=4)
    definitions_heading.paragraph_format.page_break_before = True
    widths = [1.18, 0.95, 0.63, 0.90, 0.95, 1.16]
    for group, purpose, specs in GROUPS:
        add_text(doc, group, style='Heading 1', before=7, after=2)
        add_text(doc, purpose, italic=True, after=4)
        for table_name, table_purpose in specs:
            if table_name not in tables:
                add_text(doc, f'{table_name} — SQL definition not found in the repository.', italic=True, after=3)
                continue
            fields = tables[table_name]['fields']
            if not fields:
                continue
            title = doc.add_paragraph(style='Heading 2')
            title.paragraph_format.keep_with_next = True
            format_paragraph(title, 5, 1)
            tr = title.add_run(table_name); set_run_font(tr, 12, True, color=(25,55,80))
            desc = doc.add_paragraph()
            desc.paragraph_format.keep_with_next = True
            format_paragraph(desc, 0, 2)
            dr = desc.add_run(table_purpose); set_run_font(dr, 12, False, True, (55,65,75))
            rows = []
            for col, meta in fields.items():
                data_type, length = type_and_length(meta['type'])
                rows.append((col, data_type, length, 'Yes' if meta['nullable'] else 'No', compact_default(meta['default']), key_text(meta['keys'])))
            add_table(doc, ['Column Name', 'Type', 'Length', 'Nullable', 'Default', 'Key'], rows, widths)

    enum_heading = add_text(doc, 'PostgreSQL enum values', style='Heading 1', before=8, after=3)
    enum_heading.paragraph_format.keep_with_next = True
    enum_heading.paragraph_format.page_break_before = True
    enum_description = add_text(doc, 'These enumerated types appear in the fields listed above. Text fields may also have separate CHECK constraints.', after=3)
    enum_description.paragraph_format.keep_with_next = True
    used_enum_types = set()
    for _, _, specs in GROUPS:
        for table_name, _ in specs:
            if table_name not in tables:
                continue
            for meta in tables[table_name]['fields'].values():
                type_name = re.sub(r'\s+.*$', '', meta['type'].strip()).lower()
                if type_name in ENUM_VALUES:
                    used_enum_types.add(type_name)
    enum_rows = [(name, values) for name, values in ENUM_VALUES.items() if name in used_enum_types]
    if enum_rows:
        add_table(doc, ['Enum Type', 'Allowed Values'], enum_rows, [1.5, 4.27])
    doc.core_properties.title = 'NorzAgapay Data Dictionary'
    doc.core_properties.subject = 'Database tables and fields used by NorzAgapay web workflows'
    doc.core_properties.author = 'NorzAgapay'
    doc.save(OUT)
    print(f'Created {OUT}')
    print(f'Documented tables: {sum(len(specs) for _,_,specs in GROUPS)}; schema tables available: {len(tables)}')
    print(f'Pages will be checked after rendering; total fields in dictionary: {sum(len(tables[name]["fields"]) for _,_,specs in GROUPS for name,_ in specs if name in tables)}')

if __name__ == '__main__':
    main()
