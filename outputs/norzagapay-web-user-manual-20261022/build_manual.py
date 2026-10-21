from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r'C:\Users\user\Desktop\MySystems\NorzAgapay\outputs\norzagapay-web-user-manual-20261022\NorzAgapay_User_Manual.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.45)
sec.footer_distance = Inches(0.45)

# Document-wide type and paragraph rhythm.
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.widow_control = True

for sty_name, size, before, after in [('Title', 24, 0, 6), ('Heading 1', 16, 15, 6), ('Heading 2', 13, 10, 4), ('Heading 3', 12, 8, 3)]:
    st = styles[sty_name]
    st.font.name = 'Times New Roman'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.widow_control = True

# Word's built-in Title style can carry a colored bottom border; remove it.
title_ppr = styles['Title']._element.get_or_add_pPr()
for border in list(title_ppr.findall(qn('w:pBdr'))):
    title_ppr.remove(border)

# Header and page number footer.
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run('NORZAGAPAY  |  USER MANUAL')
hr.font.name = 'Times New Roman'; hr.font.size = Pt(9); hr.font.color.rgb = RGBColor(70, 70, 70)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('NorzAgapay  |  ')
fr.font.name = 'Times New Roman'; fr.font.size = Pt(9); fr.font.color.rgb = RGBColor(75, 75, 75)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

# Reusable elements.
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, color='D9DDE3', size='5'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), size); el.set(qn('w:color'), color)

def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)

def add_table(headers, rows, widths=None, font_size=10.5, padding=95):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        if widths: cell.width = Inches(widths[i])
        shade(cell, '24476A')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=padding, bottom=padding); set_cell_border(cell)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        run = p.add_run(text); run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(font_size); run.font.color.rgb = RGBColor(255,255,255)
    repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if widths: cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i], top=padding, bottom=padding); set_cell_border(cells[i])
            if ridx % 2: shade(cells[i], 'F2F5F8')
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.02
            r = p.add_run(str(value)); r.font.name = 'Times New Roman'; r.font.size = Pt(font_size); r.font.color.rgb = RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def heading(text, level=1):
    return doc.add_heading(text, level=level)

def para(text='', bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def steps(items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.27)
        p.paragraph_format.first_line_indent = Inches(-0.27)
        p.paragraph_format.keep_together = True
        r = p.add_run(f'{len(steps.counter)+1}. '); r.bold = True
        p.add_run(item)
        steps.counter.append(1)
    steps.counter.clear()
steps.counter = []

def bullets(items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(1)
        p.add_run('– ').bold = True
        p.add_run(item)

# Title block.
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run('APPENDIX D'); r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(12)
p = doc.add_paragraph(style='Title')
ppr = p._p.get_or_add_pPr()
for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
p.add_run('NorzAgapay User Manual')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
r = p.add_run('Incident reporting and response workflows'); r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0,0,0)
para('This manual explains how residents, barangay personnel, MDRRMO responders, dispatchers, and administrators use NorzAgapay to report incidents, coordinate assistance, record field actions, and review outcomes. Follow the steps that match your account role; the menu items available to you depend on the access assigned to that account.')

heading('1 Introduction', 1)
heading('1.1 Purpose and scope', 2)
para('NorzAgapay connects incident reporting with municipal and barangay response work. The report lifecycle begins when a resident or guest submits an incident, continues as dispatchers review and route it, and ends when an authorized responder records the field assessment and resolution. Residents can follow their reports, while barangay and MDRRMO personnel use their role-specific workspaces to coordinate the response.')
para('This guide covers the NorzAgapay website’s public pages and signed-in dashboards. It describes common web workflows and the actions available to each role. It does not prescribe emergency response decisions: follow local protocols and contact emergency services directly when immediate help is required.')

heading('1.2 Users and responsibilities', 2)
add_table(['User group', 'Main tasks in NorzAgapay'], [
    ('Public visitor or guest', 'Submit an emergency or community incident without creating an account; track a guest report with its reference and mobile number.'),
    ('Resident', 'Submit and track personal reports, read community alerts, find evacuation centers, and view emergency hotlines.'),
    ('Barangay administrator or staff', 'Review local information and resolved incident records; manage the barangay roster where permitted; publish community updates.'),
    ('Barangay dispatcher or responder', 'Coordinate incidents assigned to the barangay, act on response assignments, record field findings, and request MDRRMO support when needed.'),
    ('MDRRMO dispatcher', 'Review incoming reports, classify and prioritize incidents, assign a nearby active barangay or dispatch MDRRMO responders, and monitor progress.'),
    ('MDRRMO administrator or logistics staff', 'Manage accounts, verification, officers, response units, command locations, broadcasts, and system analytics according to assigned permissions.'),
], [1.75, 4.75], 10.5)

heading('1.3 Access requirements', 2)
bullets([
    'Use the NorzAgapay website address provided by your municipality and a current browser with an internet connection.',
    'Allow location access when you want to attach your current position or let a responder workspace record GPS arrival. A map pin can be selected when current location is unavailable.',
    'Use the mobile number and password associated with your web account. Keep the SMS verification code private.',
    'Keep submitted report references and guest tracking mobile numbers so you can reopen the matching report.',
])

heading('2 NorzAgapay incident lifecycle', 1)
para('The website uses the following general cycle. A report may move through barangay coordination, MDRRMO dispatch, or both, depending on its type and the operational decision made by the dispatcher.')
add_table(['Stage', 'What happens'], [
    ('1. Report', 'A guest or resident submits an emergency incident or community concern with a location and the required report details.'),
    ('2. Review', 'MDRRMO reviews the report, checks the submitted details, and may classify it, set priority, assign a barangay, or dispatch responders.'),
    ('3. Coordinate', 'The selected barangay or MDRRMO responder sees the assigned incident in the relevant workspace and proceeds under local response procedures.'),
    ('4. Record response', 'A responder records arrival, completes the required field assessment, adds media when useful, and records transport when severe injuries require delivery.'),
    ('5. Resolve and review', 'An authorized user saves a resolution summary. The incident appears in resolved history, where permitted users can review or download its resolution document.'),
], [1.35, 5.15], 10.5)
para('Typical status labels are Pending, Responding, and Resolved. The available action buttons depend on role, assignment, and current status.')

p = heading('3 Getting started', 1)
p.paragraph_format.page_break_before = True
heading('3.1 Open the website', 2)
steps([
    'Open the NorzAgapay website address supplied by your municipality.',
    'For public reporting, choose the report option from the public landing page. No account is required for a guest report.',
    'To use a resident or operations dashboard, select Sign In and enter the account mobile number and password.',
    'After sign-in, NorzAgapay opens the dashboard associated with your role. Use the left navigation or dashboard shortcuts to open an available module.',
])
heading('3.2 Create a resident account', 2)
steps([
    'Open Create a resident account from the sign-in page.',
    'Enter your full name, mobile number, and a password of at least eight characters, then select Send verification code.',
    'Enter the six-digit SMS code sent to your number. The code expires after 10 minutes; request another code if needed.',
    'Select Verify and create account. After verification, the website opens your resident report area.',
])
para('A Barangay Account Request follows a separate approval process. Creating the account does not activate barangay dispatch access; MDRRMO must review the signed authorization information and approve access. Barangay administrators can manage their local accounts after activation.')
heading('3.3 Sign in and recover a password', 2)
steps([
    'On Sign In, enter the mobile number and password linked to your account, then select Sign In.',
    'If you forgot the password, select Forgot password and choose the account type.',
    'Enter the account mobile number and request an SMS code. Enter the code and a new password when prompted.',
    'Return to Sign In and use the new password. If the account remains inaccessible, contact the system administrator for your organization.',
])

p = heading('4 Public and resident reporting', 1)
p.paragraph_format.page_break_before = True
heading('4.1 Submit an emergency or community report', 2)
para('Guests and signed-in residents use the same report intake. When a resident is signed in, the report is linked to that resident account. Public guests can submit without signing in.')
steps([
    'Open Report an Incident. Choose Emergency incident for an urgent response or Community incident for a local hazard or concern.',
    'Choose Standard report to enter written details, or Visual report to submit camera evidence without writing incident details.',
    'For an emergency, choose the incident category and indicate when it occurred. If you choose Other, describe the emergency. For a community concern, select its concern category.',
    'Enter the required mobile number. In a standard emergency report, describe what happened and who needs help. The incident description must contain at least eight characters.',
    'Set the incident location using Use current location or Choose location on map. On the map, place the pin in Norzagaray and select Use this location. Confirm that the location indicator shows a selected point.',
    'For a visual report, capture at least one photo or video. A standard report may also include optional photo or video evidence.',
    'Review the details and select Send report to MDRRMO. Wait for the submission confirmation and keep the report reference shown on screen.',
])
para('If the location is missing, the report cannot be submitted. Visual reports require at least one photo or video. Call an emergency hotline directly if you need immediate assistance and cannot submit the form.')
heading('4.2 Track a guest report', 2)
steps([
    'Open Track a guest report from the public page or the report confirmation screen.',
    'Enter the report reference and the same mobile number used on the report, then select View report status.',
    'Review the status, incident location, report details, and available evidence.',
    'While the report remains open, use the follow-up option to send an additional emergency update. Attach photo or video evidence when it helps explain the update.',
])
para('A guest report is private to the matching reference and mobile number. Keep both values somewhere secure; do not share them publicly.')
heading('4.3 Use the resident portal', 2)
p = para('The resident menu provides the following services:')
p.paragraph_format.keep_with_next = True
add_table(['Resident area', 'How to use it'], [
    ('Home', 'Review recent reports, quick links, and current community alerts.'),
    ('Community Alerts', 'Read MDRRMO updates and posts published by the selected barangay. Refresh the feed to load recent items.'),
    ('Report an Incident', 'Choose an emergency incident or community concern and submit it using the reporting steps above.'),
    ('My Reports', 'Review your report status, response milestones, location, and resolution history.'),
    ('Evacuation Centers', 'Browse available shelters, select a center, optionally show your location, and open map directions.'),
    ('Emergency Hotlines', 'Review published MDRRMO and barangay contacts and call a number directly when immediate assistance is needed.'),
], [1.6, 4.9], 10.5)

heading('5 Barangay workflows', 1)
heading('5.1 Review the barangay overview', 2)
steps([
    'Sign in with an activated Barangay Account. The overview and menu options vary by barangay role.',
    'Open Barangay Overview to review active and resolved incident counts, pending assistance requests where available, team activity, posts, or analytics.',
    'Select an incident shortcut or open Incidents to review a report assigned to the barangay. Open the report to check its description, location, reporter information, and response history.',
    'Use Barangay Board to read MDRRMO alerts and community posts. Authorized barangay users can publish local advisories from the available board controls.',
])
heading('5.2 Coordinate an assigned incident', 2)
para('Barangay dispatchers and responders see the actions permitted by their role. A dispatcher can review an assigned report and route it to active barangay responder accounts. A barangay responder uses My Responses to work on an incident assigned to that responder. The local workflow does not require the Barangay user to select an MDRRMO unit crew.')
steps([
    'Open Incidents or My Responses and select the report assigned to your barangay or responder account.',
    'Review the incident description, location, evidence, and any MDRRMO instructions. Use View location or the location link when a map is available.',
    'Travel to the incident and record arrival. The browser may record arrival automatically when location access is enabled and the GPS fix is sufficiently accurate; otherwise, select Record arrival.',
    'Complete the field assessment after arrival. Enter Situation, People affected / urgency, Actions taken, and Risks / resources. Choose a triage outcome: Safe, Mild injuries, or Severe injuries.',
    'Use Not applicable in a section when it genuinely does not apply. Save assessment after all four sections and the triage choice are complete.',
    'Add field photos or videos when useful, then save a clear resolution summary and select Resolve incident when the response is complete.',
])
para('For Severe injuries, follow the transport panel to choose one of the nearby hospital results or search by hospital name and select a mapped result. Start transport, then select Mark delivered when the injured person reaches the destination. The incident cannot be resolved until delivery is recorded.')
heading('5.3 Request MDRRMO assistance', 2)
steps([
    'Open Assistance Requests and select an active incident assigned to your barangay.',
    'Choose Request assistance. Select the needed support, such as additional responders, supplies or resources, or equipment.',
    'Explain what is needed and how it will help, then select Submit request.',
    'Review the request status and dispatcher decision in the request history. A barangay responder can acknowledge support when assistance has been provided.',
])
para('MDRRMO dispatchers review the request and record whether to provide barangay assistance, coordinate with MDRRMO, or dismiss it. They add a note when required.')
heading('5.4 Manage barangay records and users', 2)
para('Barangay administrators can open Barangay Team to add, edit, activate, or deactivate eligible barangay accounts. The account form supports Responder, Dispatcher, and Staff roles; administrator accounts are managed through the Barangay Account Request process. Staff can view the team roster and, where granted access, review the resolved incident history. Barangay Analytics summarizes local reports and outcomes for authorized administrators.')
heading('5.5 Review resolved reports and download a document', 2)
steps([
    'Open Incidents or the resolved incident workspace available to your role, then choose the Resolved view or a resolved report.',
    'Review the field assessment, submitted media, and resolution notes.',
    'Select Download resolution document or Download resolution PDF and wait for the browser download to finish.',
    'If the download fails, refresh the resolved report and retry. If it continues to fail, give the report reference and the displayed error to the system administrator.',
])

p = heading('6 MDRRMO workflows', 1)
heading('6.1 Review reports and dispatch a response', 2)
steps([
    'Open Command Center to review the map and incident queue, or open Incidents to browse reports by status and report group.',
    'Select a pending report. Review the reporter information, description, incident time, location, and proof of incident.',
    'For a resident incident, choose Assign to barangay to select a response barangay. NorzAgapay may show the nearest active barangay; confirm the destination and enter instructions or context before assigning.',
    'To dispatch MDRRMO responders, choose Dispatch responders or Classify and dispatch. Select an incident category and priority, enter optional dispatcher notes, choose the active responder accounts shown, and select Send responders.',
    'Use the status tabs to monitor Pending, Responding, and Resolved reports. Open the report timeline to review assignment, arrival, field assessment, and closure updates.',
])
para('A dispatcher may recall or reassign an active barangay assignment, return a report to the MDRRMO queue, or mark a report invalid when those actions are available and appropriate. Review the confirmation prompt and record an explanation where requested.')
heading('6.2 Complete a field response', 2)
steps([
    'Sign in to the assigned responder account and open My Responses. Select an active assignment and review its location, notes, and MDRRMO response unit details.',
    'Use View location to open the incident map. Keep the responder workspace open while traveling if automatic GPS arrival is desired and browser location access is allowed.',
    'At the scene, confirm arrival. If GPS does not record it, select Record arrival.',
    'Complete and save the four field assessment sections and triage outcome. Add field media as needed.',
    'For severe injuries, select Transport injured. Choose one of the five nearest hospital results shown for the current responder location, or search by hospital name and select a mapped result. Start transport and mark the injured person delivered at the destination.',
    'Enter a resolution summary and select Resolve incident when the field assessment is saved and any required transport is marked delivered.',
])
para('MDRRMO responders may see MDRRMO response crew information for their assignment. Ask logistics staff to correct a missing or outdated unit roster. Barangay responders use the barangay assignment workflow and do not manage an MDRRMO crew.')
heading('6.3 Manage users and operational setup', 2)
add_table(['Module', 'Typical workflow'], [
    ('Barangay Verification', 'Review Barangay Account Requests, inspect the signed authorization information, and approve, reject, or request corrections. Activation is controlled by MDRRMO.'),
    ('User Management', 'Create or edit dashboard accounts, assign an appropriate role, update access status, and refresh the account list.'),
    ('Officers and Respond Units', 'Maintain MDRRMO officer records and response-unit details. Assign unit roles and an eligible team leader, then ensure the unit is active and dispatchable.'),
    ('Command Locations and Municipality Boundary', 'Maintain operational map locations and the configured municipal boundary for map-based work.'),
    ('Alert Broadcasts', 'Create or edit an alert, enter its message and audience details, publish it, and pin or remove the post when appropriate.'),
    ('Analytics', 'Review report counts, statuses, and response outcomes available to your administrative role.'),
], [1.75, 4.75], 10.5)

heading('7 Common messages and troubleshooting', 1)
add_table(['Issue', 'What to do'], [
    ('Sign-in is rejected', 'Check the mobile number and password. Use Forgot password to request an SMS reset code. If the account is inactive or pending, contact the account administrator.'),
    ('Resident verification code is missing or expired', 'Confirm the mobile number, check SMS reception, and select Send another code. Enter the newest six-digit code before it expires.'),
    ('Barangay dashboard access is unavailable', 'A Barangay Account Request must be reviewed and activated by MDRRMO. Refresh the request status or contact the MDRRMO administrator.'),
    ('Report cannot be submitted', 'Complete required fields, use a valid mobile number, select an incident location, and meet the minimum description or visual-evidence requirement for the chosen format.'),
    ('Current location is unavailable', 'Allow location access in the browser and try again. If GPS is still unavailable, choose the incident point directly on the map.'),
    ('Guest report cannot be found', 'Use the exact report reference and the mobile number entered on that report. These values must match.'),
    ('No responder assignment appears', 'Select Refresh in My Responses and confirm that the dispatcher assigned the incident to your account. Contact the dispatcher if it remains missing.'),
    ('Nearby hospitals did not load', 'Allow responder location access and retry the hospital lookup. Enter a hospital name and select a mapped result if the nearby list is unavailable. Contact MDRRMO if no valid destination appears.'),
    ('Resolve incident is unavailable', 'Record arrival, fill all four assessment sections, choose a triage outcome, and save the assessment. For severe injuries, start transport and mark delivery before resolving.'),
    ('Resolution document will not download', 'Confirm that the report is resolved, refresh the report details, and retry the download. If it still fails, provide the report reference and exact error message to the system administrator.'),
    ('The page shows an error or stale information', 'Use Refresh or Try again, then sign in again if needed. For a persistent issue, tell the NorzAgapay administrator or dispatcher your role, report reference if available, page, time, and exact message. Never send passwords or SMS codes.'),
], [1.85, 4.65], 10.2)

heading('8 Account security and safe use', 1)
bullets([
    'Keep your password, SMS verification codes, guest report reference, and tracking mobile number private.',
    'Enter accurate incident details and place the location pin as close as possible to the actual event.',
    'Share only information needed for the response. Treat reporter contact details and incident media as private.',
    'Confirm the report and destination before dispatching or reassigning responders. Follow local response procedures in the field.',
    'Save the assessment before resolving an incident. Do not mark severe-injury transport delivered until the injured person has reached the selected destination.',
    'Sign out after using a shared computer and close the browser tab when finished.',
])

heading('9 Quick reference', 1)
add_table(['Task', 'Where to go'], [
    ('Report or track an incident', 'Public page: Report or Track a guest report with reference and mobile number. Resident Portal: Report an Incident or My Reports.'),
    ('Read alerts or find public services', 'Resident Portal → Community Alerts, Evacuation Centers, or Emergency Hotlines'),
    ('Coordinate or respond to Barangay incidents', 'Barangay Overview, Incidents, My Responses, or Assistance Requests, as available to your role'),
    ('Dispatch or manage operations', 'Command Center, Incidents, Barangay Verification, User Management, Officers, or Respond Units'),
    ('Review a resolution or download its PDF', 'Resolved incident workspace → select report → Download resolution document'),
], [2.4, 4.1], 10.0, 40)

# Remove the optional empty paragraph after the final table to avoid a blank page.
last_paragraph = doc.paragraphs[-1]
if not last_paragraph.text.strip():
    last_paragraph._element.getparent().remove(last_paragraph._element)

# Accessibility and document properties.
doc.core_properties.title = 'NorzAgapay User Manual'
doc.core_properties.subject = 'Incident reporting and response workflows'
doc.core_properties.author = 'NorzAgapay'
doc.core_properties.keywords = 'NorzAgapay, user manual, incident reporting, response workflow'

doc.save(OUT)
print(OUT)
print(f'Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}')
